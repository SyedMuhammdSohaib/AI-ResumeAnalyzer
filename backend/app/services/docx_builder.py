"""Turn the AI's simple markdown resume into a clean Word document."""
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x8A)
MUTED = RGBColor(0x55, 0x5B, 0x6E)
BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _add_runs(paragraph, text: str, size: float | None = None, color: RGBColor | None = None):
    """Add text to a paragraph, rendering **bold** segments as bold runs."""
    for part in BOLD.split(text):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if is_bold else part)
        run.bold = is_bold or None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def _bottom_border(paragraph) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F3A8A"}.items():
        bottom.set(qn(key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _spacing(paragraph, before: float = 0, after: float = 2) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def build_docx(markdown: str) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    in_header = True  # lines between the name and the first section are contact info
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or line.startswith("```"):
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(p, after=2)
            _add_runs(p, line[2:].strip().replace("**", ""), size=22, color=ACCENT)
            p.runs[0].bold = True
        elif line.startswith("## "):
            in_header = False
            p = doc.add_paragraph()
            _spacing(p, before=10, after=4)
            _add_runs(p, line[3:].strip().replace("**", "").upper(), size=12, color=ACCENT)
            p.runs[0].bold = True
            _bottom_border(p)
        elif line.startswith("### "):
            in_header = False
            p = doc.add_paragraph()
            _spacing(p, before=6, after=1)
            _add_runs(p, line[4:].strip().replace("**", ""), size=11)
            p.runs[0].bold = True
        elif match := BULLET.match(line):
            p = doc.add_paragraph(style="List Bullet")
            _spacing(p, after=1)
            _add_runs(p, match.group(1).strip())
        elif in_header:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(p, after=1)
            _add_runs(p, line, size=9.5, color=MUTED)
        else:
            p = doc.add_paragraph()
            _spacing(p, after=3)
            _add_runs(p, line)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
