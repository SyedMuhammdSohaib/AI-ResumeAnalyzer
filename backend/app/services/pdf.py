"""Extract plain text from an uploaded resume (PDF or DOCX)."""
import zipfile
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from ..config import MAX_FILE_SIZE


class ResumeValidationError(Exception):
    """Raised when an uploaded file cannot be used as a resume."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _pdf_text(content: bytes) -> str:
    if b"%PDF-" not in content[:1024]:
        raise ResumeValidationError(
            "File does not look like a valid PDF.", status_code=415
        )

    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ResumeValidationError(
                "PDF is encrypted/password-protected. Please upload an open PDF.",
                status_code=422,
            )
        pages = reader.pages
        return "\n".join((page.extract_text() or "") for page in pages).strip()
    except ResumeValidationError:
        raise
    except Exception:
        raise ResumeValidationError(
            "Could not read the PDF. It may be corrupted.", status_code=422
        )


def _docx_text(content: bytes) -> str:
    # A .docx file is a zip archive; anything else (e.g. an old .doc renamed) is rejected.
    if not content.startswith(b"PK"):
        raise ResumeValidationError(
            "File does not look like a valid Word (.docx) document.", status_code=415
        )

    try:
        doc = Document(BytesIO(content))
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise ResumeValidationError(
            "Could not read the Word document. It may be corrupted.", status_code=422
        )

    lines = [p.text for p in doc.paragraphs]
    # Many resume templates put content in tables (e.g. two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                # Merged cells repeat the same cell object across columns.
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(line for line in lines if line.strip()).strip()


def extract_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".doc"):
        raise ResumeValidationError(
            "Old Word (.doc) files are not supported. Save it as .docx or PDF and try again.",
            status_code=415,
        )
    if not name.endswith((".pdf", ".docx")):
        raise ResumeValidationError(
            "Invalid file type. Please upload a PDF or Word (.docx) resume.", status_code=415
        )

    if not content:
        raise ResumeValidationError(
            "The uploaded file is empty.", status_code=422
        )

    if len(content) > MAX_FILE_SIZE:
        raise ResumeValidationError(
            "File is too large. Maximum size is 5 MB.", status_code=413
        )

    text = _pdf_text(content) if name.endswith(".pdf") else _docx_text(content)

    if not text:
        raise ResumeValidationError(
            "No extractable text found in the file (it may be a scanned image).",
            status_code=422,
        )

    return text
